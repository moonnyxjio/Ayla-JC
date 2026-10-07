"""Export everything a teacher should proofread into one Excel file.

    python workbook/export_review.py     # writes workbook/dist/파닉스워크북_감수용.xlsx

Sheets: 사용법 · 예문 · 단어 뜻·그림 · 음가 · 발음 팁 · 목차 점검
Teachers edit only the yellow columns; send the file back and the edits can be merged into content.py / eumga.py.
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_workbook as WB  # noqa: E402
import eumga as E  # noqa: E402
from content import EMOJI, GLOSS, SENTENCES, SUPPLEMENT, TIPS  # noqa: E402

OUT = os.path.join(HERE, "dist", "파닉스워크북_감수용.xlsx")
FONT = "맑은 고딕"
HEAD = PatternFill("solid", fgColor="1F2A44")
EDIT = PatternFill("solid", fgColor="FFF4C2")
THIN = Side(style="thin", color="D5DAE3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet(wb, title, headers, widths, rows, edit_cols=(), note=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD if headers[c - 1] not in edit_cols else PatternFill("solid", fgColor="B7791F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BOX
            if headers[c - 1] in edit_cols:
                cell.fill = EDIT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 30
    if note:
        ws.cell(ws.max_row + 2, 1, note).font = Font(name=FONT, size=9, italic=True, color="6B7487")
    return ws


def build():
    rows = WB.build_lesson_data(WB.load_rows())
    wb = Workbook()
    wb.remove(wb.active)

    # 사용법
    ws = wb.create_sheet("사용법")
    lines = [
        ("파닉스 워크북 감수용 파일", True),
        ("", False),
        ("노란 칸만 고쳐 주세요. 흰 칸은 현재 교재에 들어간 내용입니다.", False),
        ("고친 파일을 보내 주시면 교재 전체(통합본 · 1~5권 · 흑백판)에 한 번에 반영됩니다.", False),
        ("", False),
        ("시트 안내", True),
        ("예문 — B쪽 ⑧ 문장 완성 · C쪽 ⑫ 문장 쓰기 · 다음 단계 ⑨ 문장 받아쓰기에 쓰이는 84문장", False),
        ("단어 뜻·그림 — 읽기 카드 · 그림 연결 · 그림 보고 쓰기 · 빈칸 힌트에 쓰이는 뜻과 그림(이모지)", False),
        ("음가 — 음가 쓰기 페이지와 Phonics Test의 정답(칸별 음가 · 합쳐 읽기 · 음절 나누기)", False),
        ("발음 팁 — A쪽 맨 위 💡 소리 팁 42개", False),
        ("목차 점검 — 원본 목차(엑셀)의 단어 수 기재값과 실제 단어 수 비교", False),
        ("", False),
        ("음가 기준 (현재)", True),
        ("짧은 모음 a=ㅐ e=ㅔ i=ㅣ o=ㅏ u=ㅓ · 긴 모음 a=에이 i=아이 u=유 e=이", False),
        ("긴 o: 끝소리면 오우(go 고우, snow 스노우), 뒤에 자음이 붙으면 오(boat 보트, phone 폰)", False),
        ("au/aw = 아우 (학원 기준) · r 소리 모음은 칸에 r 표시(ar = ㅏr), 합쳐 읽기에서는 생략(car 카)", False),
        ("받침: 짧은 모음 뒤 끝소리 t·k·p → ㅅ·ㄱ·ㅂ (cat 캣), n·m·l·ng → 받침, 그 밖은 ㅡ (dog 다그)", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11)
    ws.cell(3, 1).fill = EDIT
    ws.column_dimensions["A"].width = 110

    # 예문
    data = []
    for r in rows:
        for k, (s, a) in enumerate(SENTENCES[r["page"]], 1):
            uses = "B⑧ 문장 완성 · C⑫ 문장 쓰기" + (" · 다음 단계 ⑨ 문장 받아쓰기" if k == 2 and r["page"] < 42 else "")
            data.append([r["page"], r["set"], r["sound"], k, s, a, s.replace("___", a), uses, None, None])
    sheet(wb, "예문", ["단계", "세트", "소리", "번호", "빈칸 문장", "정답", "완성 문장", "쓰이는 곳", "수정 문장", "메모"],
          [6, 7, 16, 6, 34, 10, 34, 30, 34, 24], data, edit_cols=("수정 문장", "메모"),
          note="빈칸은 ___ 로 표시해 주세요. 정답 단어는 그 단계의 단어 중에서 골라야 합니다.")

    # 단어 뜻·그림
    data = []
    for r in rows:
        p = r["page"]
        used = {}
        for w, _ in r["read"]:
            used.setdefault(w, []).append("읽기")
        for w in r["practice"]:
            used.setdefault(w, []).append("B 따라쓰기·빈칸")
        for w in r["dictation"]:
            used.setdefault(w, []).append("B 받아쓰기")
        for w in r["match"][0]:
            used.setdefault(w, []).append("B 그림·뜻 연결")
        for w in r["pictures"]:
            used.setdefault(w, []).append("C 그림 보고 쓰기")
        for w, u in used.items():
            lw = w.lower()
            data.append([p, r["set"], w, "보충" if w in SUPPLEMENT.get(p, []) else "", GLOSS.get(lw, ""),
                         EMOJI.get(lw, ""), " · ".join(dict.fromkeys(u)), None, None, None])
    sheet(wb, "단어 뜻·그림", ["단계", "세트", "단어", "보충", "뜻", "그림", "쓰이는 곳", "수정 뜻", "수정 그림", "메모"],
          [6, 7, 13, 6, 18, 7, 40, 18, 10, 24], data, edit_cols=("수정 뜻", "수정 그림", "메모"),
          note="뜻이 비어 있는 단어는 현재 교재에서 뜻 없이 쓰입니다. 뜻을 적어 주시면 그림·뜻 연결과 힌트에 쓸 수 있습니다.")

    # 음가
    data = []
    for r in rows:
        p = r["page"]
        for w in dict.fromkeys([w for w, _ in r["read"]] + r["dict"]):
            cells = E.eumga(p, w)
            exc = "예외(직접 입력)" if w.lower() in E.CELL_OVERRIDE or w.lower() in E.READING_OVERRIDE else ""
            data.append([p, r["set"], w, " | ".join(g for g, _ in cells), " | ".join(j for _, j in cells),
                         E.reading(p, w), " | ".join(E.syllables(p, w)), exc, None, None, None])
    sheet(wb, "음가", ["단계", "세트", "단어", "철자 나누기", "칸별 음가", "합쳐 읽기", "음절", "처리", "수정 음가", "수정 읽기", "메모"],
          [6, 7, 12, 18, 20, 12, 14, 14, 20, 12, 24], data, edit_cols=("수정 음가", "수정 읽기", "메모"),
          note="수정 음가는 칸별로 | 를 넣어 적어 주세요. 예) ㅋ | 에이 | ㅋ | ×")

    # 발음 팁
    data = [[r["page"], r["set"], r["sound"], TIPS[r["page"]], None, None] for r in rows]
    sheet(wb, "발음 팁", ["단계", "세트", "소리", "현재 팁", "수정 팁", "메모"], [6, 7, 18, 70, 70, 24], data,
          edit_cols=("수정 팁", "메모"))

    # 목차 점검 (formulas count the words actually listed)
    import openpyxl
    src = openpyxl.load_workbook(WB.SRC, data_only=True).worksheets[0]
    data = []
    for i, row in enumerate(src.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        p = int(row[0])
        e = f"E{len(data) + 2}"
        cnt = f'=LEN({e})-LEN(SUBSTITUTE({e},",",""))+LEN({e})-LEN(SUBSTITUTE({e},"·",""))+1'
        n = len(data) + 2
        data.append([p, row[1], row[2], row[3], row[4], row[5], row[6], cnt, f"=G{n}-H{n}",
                     len(SUPPLEMENT.get(p, [])), None])
    ws = sheet(wb, "목차 점검", ["페이지", "세트", "소리", "묶음 기준", "읽기/예시 단어", "받아쓰기 단어", "단어수(기재)",
                               "실제 단어 수", "차이", "보충 단어 수", "메모"],
               [7, 7, 18, 26, 50, 34, 10, 10, 8, 10, 24], data, edit_cols=("메모",),
               note="실제 단어 수·차이는 수식입니다. 차이가 0이 아니면 원본 G열 기재값이 실제 단어 수와 다릅니다. "
                    "보충 단어 수는 교재가 원본에 없는 소리를 채우려고 추가한 단어 수(출처: content.py SUPPLEMENT)입니다.")
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add(f"I2:I{len(data) + 1}",
                                  CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor="FBD5CF")))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT, file=sys.stderr)


if __name__ == "__main__":
    build()
