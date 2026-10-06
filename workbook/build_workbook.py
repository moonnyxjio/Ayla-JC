"""Build a printable phonics workbook (HTML -> PDF) from phonics_toc.xlsx + content.py.

Usage:
    python workbook/build_workbook.py      # writes workbook/phonics_workbook.html + .pdf

Every TOC row becomes a three-page lesson:
  A 배우기  — 소리 팁(+QR), 읽기, 소리 블렌딩, 소리 찾기/분류, 지난 소리 복습
  B 연습    — 따라 쓰기, 빈칸 채우기, 그림·뜻 연결, 문장 완성, 받아쓰기(+QR), 자기 점검
  C 쓰기    — 그림 보고 쓰기, 보고 쓰기, 문장 따라 쓰기·옮겨 쓰기, 나만의 문장
Each Set ends with two review pages, and the answer key is at the back.
"""
import html
import os
import random
import re
import subprocess
import sys
from urllib.parse import quote, quote_plus

import openpyxl
import segno

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from content import EMOJI, FRAMES, GLOSS, SENTENCES, SORT, SUPPLEMENT, SYLLABLES, TIPS  # noqa: E402

SRC = os.path.join(HERE, "phonics_toc.xlsx")
CSS_FILE = os.path.join(HERE, "workbook.css")
OUT_HTML = os.path.join(HERE, "phonics_workbook.html")
OUT_PDF = os.path.join(HERE, "phonics_workbook.pdf")
CHROME = os.environ.get("CHROME", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome")

SETS = {
    "Set 1": ("단모음 · 단자음", "#E4572E"),
    "Set 2": ("이중자 · 자음군", "#2E86AB"),
    "Set 3": ("장모음 · 모음팀", "#3BA55C"),
    "Set 4": ("r통제 · 이중모음 · 특수규칙", "#8E5BD0"),
    "Set 5": ("음절 · 어미 · 접사", "#E09F1F"),
}
SET_END = (8, 16, 26, 34, 42)

# Per-page highlight rule. A regex with groups highlights only the groups;
# otherwise every whole match is highlighted.
RULES = {
    1: r"^[aeiou]", 2: r"[tnm]", 3: r"[spd]", 4: r"[gcr]", 5: r"[bhl]", 6: r"[fkw]",
    7: r"[jvy]", 8: r"qu|[zx]", 9: r"sh|ch", 10: r"th", 11: r"tch", 12: r"ck|ng",
    13: r"wh|ph", 14: r"ss|ll|ff|zz",
    15: r"^(?:bl|cl|fl|gl|pl|sl|br|cr|dr|fr|gr|pr|tr|sc|sk|sm|sn|sp|st|sw)",
    16: r"(?:nd|nt|nk|mp|st|lt)$|^(?:spl|str|scr|spr|thr|shr|squ)",
    17: r"[aeiou]$", 18: r"([aeiou])[^aeiou](e)$", 19: r"ai|ay", 20: r"ee|ea",
    21: r"oa|ow", 22: r"igh|ie|y$", 23: r"y$", 24: r"oo|ew|ue|ui", 25: r"oo",
    26: r"eigh", 27: r"ar", 28: r"ore|or", 29: r"er|ir|ur", 30: r"oi|oy", 31: r"ou|ow",
    32: r"au|aw", 33: r"[cg](?=[eiy])", 34: r"kn|wr|mb|gn", 36: r"le$",
    37: r"(?:es|ed|ing|s)$", 38: r"(?:est|er|ly)$", 39: r"^(?:un|re|dis)",
    40: r"^(?:mis|pre)", 41: r"(?:ful|less|ness)$", 42: r"(?:tion|sion|ture)$",
}

# Page 35 (schwa): the weak vowel can't be found by spelling, so mark it by index.
SCHWA = {
    "about": [0], "sofa": [3], "zebra": [4], "banana": [1, 5],
    "pencil": [4], "problem": [5], "animal": [4], "lemon": [3],
}

GRAPHEMES = sorted([
    "eigh", "tch", "igh", "ore", "ss", "ll", "ff", "zz", "sh", "ch", "th", "ck", "ng", "wh",
    "ph", "qu", "ai", "ay", "ee", "ea", "oa", "ow", "ie", "oo", "ew", "ue", "ui", "ar", "or",
    "er", "ir", "ur", "oi", "oy", "ou", "au", "aw", "kn", "wr", "mb", "gn",
    "pp", "nn", "tt", "dd", "gg", "rr", "mm", "bb",
], key=len, reverse=True)


# ---------------------------------------------------------------- word logic
def spans(page, word):
    """Return [(start, end)] character spans to highlight in word."""
    w = word.lower()
    if page == 35:
        return [(i, i + 1) for i in SCHWA.get(w, [])]
    rule = RULES.get(page)
    if not rule:
        return []
    out = []
    for m in re.finditer(rule, w):
        if m.re.groups:
            out += [m.span(g) for g in range(1, m.re.groups + 1) if m.group(g)]
        else:
            out.append(m.span())
    return out


def render(word, sp, mode="hi"):
    """'hi' colors spans, 'blank' turns each spanned letter into a box."""
    res, i = [], 0
    for s, e in sorted(sp):
        res.append(html.escape(word[i:s]))
        seg = word[s:e]
        if mode == "hi":
            res.append(f'<b class="hi">{html.escape(seg)}</b>')
        else:
            res.append('<span class="box"></span>' * len(seg))
        i = e
    res.append(html.escape(word[i:]))
    return "".join(res)


def classify(page, word):
    """Column label for the 소리 분류 activity."""
    sp = spans(page, word)
    if not sp:
        return None
    s, e = sp[0]
    seg = word[s:e]
    if page == 15:
        return "s 블렌드" if seg[0] == "s" else ("l 블렌드" if seg[1] == "l" else "r 블렌드")
    if page == 16:
        return "3글자 자음군" if s == 0 else "끝 자음군"
    if page == 18:
        return seg + "_e"
    if page == 33:
        return "soft " + seg
    if page in (37, 38, 41, 42):
        return "-" + seg
    if page in (39, 40):
        return seg + "-"
    return seg


def chunks(page, word):
    """Split a word into sound chunks for 소리 블렌딩. Returns [(text, silent)]."""
    w = word.lower()
    if page >= 35 and w in SYLLABLES:
        return [(c, False) for c in SYLLABLES[w]]
    if page >= 37:
        sp = spans(page, w)
        if sp:
            s, e = sp[0]
            return [(c, False) for c in (w[:s], w[s:e], w[e:]) if c]
        return [(w, False)]
    out, i = [], 0
    # final e after a consonant is silent (cake, apple), except in short words like me/she
    magic = len(w) >= 4 and w.endswith("e") and w[-2] not in "aeiou"
    while i < len(w):
        if magic and i == len(w) - 1:
            out.append(("e", True))
            break
        for g in GRAPHEMES:
            if w.startswith(g, i):
                out.append((g, False))
                i += len(g)
                break
        else:
            out.append((w[i], False))
            i += 1
    return out


def parse_words(text):
    """Return [(word, tag)] from a TOC cell. Handles 'a: ant, apple · e: ...' and 'closed(cat)'."""
    words = []
    for chunk in re.split(r"\s*·\s*", str(text)):
        chunk = re.sub(r"^\s*\w\s*:\s*", "", chunk)
        for item in chunk.split(","):
            item = item.strip()
            if not item:
                continue
            m = re.match(r"([\w -]+?)\s*\((\w+)\)", item)
            w, tag = (m.group(2), m.group(1)) if m else (item, "")
            if w not in [x for x, _ in words]:
                words.append((w, tag))
    return words


def load_rows():
    ws = openpyxl.load_workbook(SRC, data_only=True).worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        page, set_, sound, rule, read, dict_ = r[:6]
        page = int(page)
        words = parse_words(read)
        for w in SUPPLEMENT.get(page, []):
            if w not in [x for x, _ in words]:
                words.append((w, "★"))
        rows.append(dict(
            page=page, set=set_, sound=sound, rule=rule, read=words[:16],
            dict=[w for w, _ in parse_words(dict_)], color=SETS[set_][1],
        ))
    return rows


# ---------------------------------------------------------------- activities
def build_lesson_data(rows):
    """Pre-compute every activity (and its answers) so pages and the key agree."""
    by_page = {r["page"]: r for r in rows}
    for r in rows:
        p = r["page"]
        rng = random.Random(p * 7919)
        read = [w for w, _ in r["read"]]

        # blend: read words not already used in the writing activities
        pool = [w for w in read if w not in r["dict"]]
        r["blend"] = (pool + [w for w in read if w not in pool])[:4]

        # sort (multi-spelling pages) or find-the-sound O/X (others)
        if p in SORT:
            cols = SORT[p]
            bucket = {c: [w for w in read if classify(p, w) == c] for c in cols}
            picked = []
            while len(picked) < 10 and any(bucket.values()):
                for c in cols:
                    if bucket[c] and len(picked) < 10:
                        picked.append(bucket[c].pop(0))
            rng.shuffle(picked)
            r["sort"] = (cols, picked, {c: [w for w in picked if classify(p, w) == c] for c in cols})
        else:
            today = [w for w in read if spans(p, w)][:6]
            review = []
            for q in range(p - 1, 0, -1):
                for w, _ in by_page[q]["read"]:
                    if not spans(p, w) and w not in review and w not in read and len(review) < 4:
                        review.append(w)
                if len(review) >= 4:
                    break
            items = today + review
            rng.shuffle(items)
            r["find"] = (items, today)

        # spaced review: 2 words each from 1, 3 and 7 lessons back
        r["review"] = []
        for back in (1, 3, 7):
            q = by_page.get(p - back)
            if q:
                ws = [w for w, _ in q["read"]]
                a, b = ws[back % len(ws)], ws[(back * 3 + 1) % len(ws)]
                if a == b:
                    b = ws[(back * 3 + 2) % len(ws)]
                r["review"].append((q, a, b))

        # match: 4 words with glosses, dictation words first
        match = []
        for w in r["dict"] + read:
            if w.lower() in GLOSS and w.lower() not in [m.lower() for m in match]:
                match.append(w)
        match = match[:4]
        meanings = [GLOSS[w.lower()] for w in match]
        rng.shuffle(meanings)
        r["match"] = (match, meanings)

        # sentences + word bank (two answers + one distractor)
        sents = SENTENCES[p]
        answers = [a.lower() for _, a in sents]
        distract = [w for w in read if w.lower() not in answers]
        bank = [a for _, a in sents] + distract[:1]
        rng.shuffle(bank)
        r["sent"] = (sents, bank)
        r["dict_sentence"] = sents[0][0].replace("___", sents[0][1])
        r["full_sentences"] = [t.replace("___", a) for t, a in sents]

        # C page: picture writing (emoji first, then meaning-only), then copy-writing words
        allw = []
        for w in r["dict"] + read:
            if w.lower() not in [x.lower() for x in allw]:
                allw.append(w)
        pics = [w for w in allw if w.lower() in EMOJI][:9]
        pics += [w for w in allw if w not in pics and w.lower() in GLOSS][:9 - len(pics)]
        r["pictures"] = pics[:6]
        rest = [w for w in r["blend"] + read if w not in pics and w not in r["dict"]]
        r["copy"] = list(dict.fromkeys(rest + [w for w in read if w not in rest]))[:3]
        r["frames"] = [FRAMES[p % len(FRAMES)], FRAMES[(p + 2) % len(FRAMES)]]
    return rows


def word_search(words, seed, size=10):
    rng = random.Random(seed)
    grid = [[None] * size for _ in range(size)]
    placed = []
    for w in words:
        W = w.upper()
        for _ in range(400):
            d = rng.choice([(0, 1), (1, 0)])
            r0 = rng.randrange(size - (len(W) - 1) * d[0])
            c0 = rng.randrange(size - (len(W) - 1) * d[1])
            cells = [(r0 + i * d[0], c0 + i * d[1]) for i in range(len(W))]
            if all(grid[a][b] in (None, W[i]) for i, (a, b) in enumerate(cells)):
                for i, (a, b) in enumerate(cells):
                    grid[a][b] = W[i]
                placed.append((w, cells))
                break
    sol = {cell for _, cells in placed for cell in cells}
    grid = [[ch or rng.choice("ABCDEFGHIKLMNOPRSTUVWY") for ch in row] for row in grid]
    return grid, placed, sol


def build_reviews(rows):
    reviews = {}
    for s in SETS:
        n = int(s[-1])
        rng = random.Random(n * 977)
        srows = [r for r in rows if r["set"] == s]

        per = 2 if len(srows) <= 10 else 1
        mixed = [(r, w) for r in srows for w, _ in r["read"][2:2 + per]][:20]

        def first_seg(page, w):
            sp = spans(page, w)
            return (sp[0], w[sp[0][0]:sp[0][1]]) if sp else (None, None)

        segs = []
        for o in srows:
            for ow in o["dict"]:
                _, seg = first_seg(o["page"], ow)
                if seg and o["page"] not in (35, 36) and seg not in segs:
                    segs.append(seg)
        choice = []
        rr = [r for r in srows if r["page"] not in (35, 36)]
        rng.shuffle(rr)
        for r in rr:
            if len(choice) == 8:
                break
            cands = [w for w in r["dict"] if spans(r["page"], w)]
            w = next((w for w in cands if w.lower() in GLOSS), cands[0] if cands else None)
            if not w:
                continue
            sp, right = first_seg(r["page"], w)
            others = [x for x in segs if x != right]
            wrong = rng.choice(others) if others else "?"
            opts = [right, wrong]
            rng.shuffle(opts)
            choice.append((r, w, sp, opts))

        pool = []
        for r in srows:
            for w in r["dict"]:
                if 3 <= len(w) <= 7 and w.isalpha() and w not in pool:
                    pool.append(w)
        rng.shuffle(pool)
        grid, placed, sol = word_search(pool[:8], n * 31)
        used = {w for w, _ in placed}
        reviews[s] = dict(rows=srows, mixed=mixed, choice=choice, grid=grid, placed=placed,
                          sol=sol, dict=[w for w in pool if w not in used][:8])
    return reviews


# ---------------------------------------------------------------- html parts
def sound_size(text, big=30):
    n = len(text)
    if n <= 14:
        return f"{big}pt"
    if n <= 26:
        return f"{big * .73:.0f}pt"
    return f"{big * .5:.0f}pt" if n <= 50 else f"{big * .4:.0f}pt"


def header(r, part):
    label = {"A": "A · 배우기", "B": "B · 연습", "C": "C · 쓰기"}[part]
    return f"""
  <div class="hd">
    <div class="tag"><div class="set">{html.escape(r['set'])}</div><div class="no">{r['page']}</div><div class="part">{label}</div></div>
    <div class="title"><div class="sound en" style="font-size:{sound_size(r['sound'], 30 if part == 'A' else 20)}">{html.escape(r['sound'])}</div>
      <div class="rule">{html.escape(r['rule'])}</div></div>
    <div class="meta"><div class="ln">날짜&nbsp;&nbsp;&nbsp;&nbsp;/</div><div>오늘의 별</div><div class="stars">★★★</div></div>
  </div>"""


def footer(left, right):
    return f"""
  <div class="ft"><span>{left}</span><span class="sign">확인 ________</span><span>{right}</span></div>"""


def lesson_footer(r, part):
    return footer(f"Phonics Workbook · {r['set']} {SETS[r['set']][0]}", f"{r['page']}{part} / 42")


def sec(n, title, hint, body, cls=""):
    return f"""
  <div class="sec {cls}"><h2><span class="n">{n}</span>{title} <small>{hint}</small></h2>{body}</div>"""


def pic(w):
    return EMOJI.get(w.lower(), "")


def qr(url, label, sub):
    svg = segno.make(url, error="l").svg_inline(scale=1, border=0, dark="#1d2433")
    return f'<div class="qr"><div class="qimg">{svg}</div><div class="qt"><b>{label}</b><span>{sub}</span></div></div>'


def listen_url(words):
    return "https://translate.google.com/?sl=en&tl=ko&op=translate&text=" + quote("\n".join(words))


VIDEO_QUERY = {
    15: "phonics beginning blends l r s", 16: "phonics ending blends", 17: "phonics open syllable long vowel",
    18: "phonics magic e silent e", 35: "schwa sound for kids", 36: "six syllable types for kids",
    37: "phonics suffix s es ed ing", 38: "suffix er est ly for kids", 39: "prefix un re dis for kids",
    40: "prefix mis pre for kids", 41: "suffix ful less ness for kids", 42: "tion sion ture phonics",
}


def video_url(sound, page=0):
    if page in VIDEO_QUERY:
        return "https://www.youtube.com/results?search_query=" + quote_plus(VIDEO_QUERY[page])
    q = re.sub(r"[^\x00-\x7f]", " ", sound)
    q = re.sub(r"\(\s*\)|[·()]", " ", q)
    return "https://www.youtube.com/results?search_query=" + quote_plus(" ".join(f"phonics {q} sound".split()))


def guide_row(h, text="", cls="", seps=(), left=0):
    """A 4-line handwriting guide h mm tall. text is drawn so its x-height fills the middle band."""
    word = (f'<div class="word {cls}" style="font-size:{40 * h / 15:.1f}pt;top:{-2.3 * h / 15:.2f}mm;left:{left + 2}mm">'
            f'{html.escape(text)}</div>') if text else ""
    sep = "".join(f'<div class="sep" style="left:{x:.0f}mm"></div>' for x in seps)
    return (f'<div class="row" style="height:{h}mm"><div class="g g1"></div>'
            f'<div class="g g2" style="top:{h / 3:.2f}mm"></div><div class="g g3" style="top:{2 * h / 3:.2f}mm"></div>'
            f'<div class="g g4" style="top:{h}mm"></div>{word}{sep}</div>')


def page_a(r):
    p = r["page"]
    tip = (f'<div class="tiprow"><div class="tip"><span class="ic">💡</span><div><b>소리 팁</b> {html.escape(TIPS[p])}</div></div>'
           f'{qr(listen_url([w for w, _ in r["read"]]), "단어 듣기", "찍으면 발음이 나와요")}'
           f'{qr(video_url(r["sound"], p), "소리 영상", "파닉스 노래·영상")}</div>')

    cells = []
    for w, tag in r["read"]:
        lab = ""
        if tag == "★":
            lab = '<span class="star">★</span>'
        elif tag:
            lab = f'<div class="lab">{html.escape(tag)}</div>'
        em = f'<span class="em">{pic(w)}</span>' if pic(w) else ""
        cells.append(f'<div class="w">{em}{lab}<div class="t en">{render(w, spans(p, w))}</div><div class="ck">○○○</div></div>')
    legend = '<div class="legend">★ 보충 단어 — 원본 목차에 없던 소리를 채운 단어예요.</div>' \
        if any(t == "★" for _, t in r["read"]) else ""
    read = sec(1, "소리 내어 읽기", "색깔 글자 소리에 집중! 한 번 읽을 때마다 ○ 하나씩 칠해요.",
               f'<div class="read">{"".join(cells)}</div>{legend}')

    bl = []
    for w in r["blend"]:
        parts = "".join(f'<span class="chunk{" silent" if silent else ""}">{html.escape(c)}</span>'
                        for c, silent in chunks(p, w))
        bl.append(f'<div class="bl"><div class="parts en">{parts}</div><div class="arr">→</div><div class="wl"></div></div>')
    hint = ("칸마다 소리를 하나씩 말하고, 빠르게 합쳐 읽은 뒤 단어를 써요." if p < 35
            else "덩어리(음절·어미)를 하나씩 읽고, 합쳐 읽은 뒤 단어를 써요.")
    if p == 18:
        hint += " 회색 e는 소리 없는 Magic e!"
    blend = sec(2, "소리 블렌딩", hint, f'<div class="blend">{"".join(bl)}</div>')

    if "sort" in r:
        cols, picked, _ = r["sort"]
        bank = " ".join(f'<span class="chipw en">{html.escape(w)}</span>' for w in picked)
        rows_per = max(2, -(-len(picked) // len(cols)) + 1)
        colhtml = "".join(f'<div class="col"><div class="ch en">{html.escape(c)}</div>'
                          + '<div class="cl"></div>' * rows_per + '</div>' for c in cols)
        act = sec(3, "소리 분류", "단어를 읽고 같은 철자끼리 표에 옮겨 써요.",
                  f'<div class="bank2">{bank}</div>'
                  f'<div class="sort" style="grid-template-columns:repeat({len(cols)},1fr)">{colhtml}</div>')
    else:
        items, _ = r["find"]
        cells = "".join(f'<div class="fx"><span class="en">{html.escape(w)}</span><span class="ox">○ ✕</span></div>'
                        for w in items)
        act = sec(3, "소리 찾기", f"오늘의 소리 <b class='hi en'>{html.escape(r['sound'])}</b> 가 있으면 ○, 없으면 ✕ 에 표시해요.",
                  f'<div class="find">{cells}</div>')

    rv = "".join(
        f'<div class="rv" style="--c:{q["color"]}"><span class="pg">{q["page"]}쪽</span>'
        f'<span class="en">{render(a, spans(q["page"], a))}</span>'
        f'<span class="en">{render(b, spans(q["page"], b))}</span><span class="ck">○○</span></div>'
        for q, a, b in r["review"])
    if rv:
        review = sec(4, "지난 소리 복습", "1·3·7쪽 전에 배운 단어를 다시 읽어요. 간격을 두고 다시 보면 오래 기억해요!",
                     f'<div class="review">{rv}</div>')
    else:
        review = sec(4, "소리로 말하기", "알파벳 이름(에이, 이…)이 아니라 '소리'로 말해 봐요.",
                     '<div class="review"><div class="rv"><span class="en">a /æ/ · e /ɛ/ · i /ɪ/ · o /ɒ/ · u /ʌ/</span>'
                     '<span class="ck">○○○</span></div></div>')

    return f"""
<section class="page" style="--c:{r['color']}">{header(r, 'A')}{tip}{read}{blend}{act}{review}{lesson_footer(r, 'A')}
</section>"""


def page_b(r):
    p = r["page"]
    rows = []
    for w in r["dict"][:5]:
        width = max(46, len(w) * 8.6 + 12)
        rows.append(f'<div class="row"><div class="g g1"></div><div class="g g2"></div><div class="g g3"></div>'
                    f'<div class="g g4"></div><div class="word">{html.escape(w)}</div>'
                    f'<div class="sep" style="left:{width:.0f}mm"></div></div>')
    trace = sec(5, "따라 쓰기", "회색 글자를 따라 쓰고, 점선 오른쪽에 두 번 더 써요.", "".join(rows), "trace")

    fills = []
    for i, w in enumerate(r["dict"][:5], 1):
        sp = spans(p, w)
        t = render(w, sp, "blank") if sp else html.escape(w)
        fills.append(f'<div class="f"><div class="t en">{t}</div><div class="n">{i}</div></div>')
    fill = sec(6, "빈칸 채우기", f"힌트 소리: <b class='hi en'>{html.escape(r['sound'])}</b>",
               f'<div class="fill">{"".join(fills)}</div>')

    words, meanings = r["match"]
    left = "".join(f'<div class="mi"><span class="en">{html.escape(w)}</span><span class="dot">●</span></div>' for w in words)
    emo = {GLOSS[w.lower()]: pic(w) for w in words}
    right = "".join(f'<div class="mi r"><span class="dot">●</span><span class="pe">{emo.get(m, "")}</span>'
                    f'<span>{html.escape(m)}</span></div>' for m in meanings)
    sents, bank = r["sent"]
    sl = "".join(
        f'<div class="sn en"><span class="num">{i}.</span><span>{html.escape(s).replace("___", "<span class=blank></span>")}</span>'
        f'<span class="ck">○○</span></div>' for i, (s, _) in enumerate(sents, 1))
    bankh = " ".join(f'<span class="chipw en">{html.escape(w)}</span>' for w in bank)
    duo = f"""
  <div class="duo">
    {sec(7, "그림·뜻 연결", "단어와 그림·뜻을 선으로 이어요.", f'<div class="match"><div>{left}</div><div>{right}</div></div>')}
    {sec(8, "문장 완성", "빈칸에 단어를 쓰고 두 번 읽어요.", f'<div class="bank2">{bankh}</div>{sl}')}
  </div>"""

    lines = "".join(f'<div class="d">{i}.</div>' for i in range(1, 6))
    dq = qr(listen_url(r["dict"][:5] + [r["dict_sentence"]]), "혼자 받아쓰기", "QR로 듣고 써요")
    dictation = sec(9, "받아쓰기", "단어 5개와 문장 1개를 듣고 써요. (어른이 불러 주거나 QR로 들어요)",
                    f'<div class="dwrap"><div class="dict">{lines}<div class="d full">문장.</div></div>{dq}</div>')

    check = """
  <div class="self"><b>자기 점검</b>
    <span>오늘의 소리를 <b>읽을</b> 수 있어요 😀 🙂 😐</span>
    <span><b>쓸</b> 수 있어요 😀 🙂 😐</span>
    <span class="redo">틀린 단어 다시 쓰기</span></div>"""
    return f"""
<section class="page" style="--c:{r['color']}">{header(r, 'B')}{trace}{fill}{duo}{dictation}{check}{lesson_footer(r, 'B')}
</section>"""


def page_c(r):
    p = r["page"]
    cells = []
    for w in r["pictures"]:
        e = pic(w)
        top = (f'<span class="pe">{e}</span><span class="gl">{html.escape(GLOSS.get(w.lower(), ""))}</span>' if e
               else f'<span class="gl big">{html.escape(GLOSS.get(w.lower(), ""))}</span>')
        cells.append(f'<div class="pc"><div class="ptop">{top}<span class="cnt">{"□" * len(w)}</span></div>'
                     f'{guide_row(12)}</div>')
    s10 = sec(10, "그림 보고 쓰기", "그림과 뜻을 보고 영어 단어를 써요. □는 글자 수예요. 막히면 A쪽 읽기 칸을 봐요.",
              f'<div class="pics">{"".join(cells)}</div>')

    rows = []
    for w in r["copy"]:
        n = 3 if len(w) <= 5 else 2
        width = 160 / n
        rows.append(f'<div class="cp"><div class="model en">{render(w, spans(p, w))}</div>'
                    f'<div class="cg">{guide_row(13, seps=[width * i for i in range(1, n)])}</div></div>')
    s11 = sec(11, "보고 쓰기", "왼쪽 단어를 보고, 소리를 말하면서 칸마다 한 번씩 써요.", "".join(rows))

    srows = []
    for t in r["full_sentences"]:
        srows.append(f'{guide_row(11, t, "sent")}<div class="gap"></div>{guide_row(11)}')
    s12 = sec(12, "문장 따라 쓰기 · 옮겨 쓰기", "윗줄은 회색 글자를 따라 쓰고, 아랫줄에는 혼자 써요.",
              '<div class="gap2"></div>'.join(srows))

    fr = "".join(
        f'<div class="frm"><span class="en fr">{html.escape(f).replace("___", "<span class=blank></span>")}</span>'
        f'{guide_row(11)}</div>' for f in r["frames"])
    s13 = sec(13, "나만의 문장", "오늘 배운 단어를 빈칸에 넣어 문장을 완성해서 써요.", fr)
    redo = "".join(f'<div class="cp"><div class="model blankm"></div><div class="cg">{guide_row(13, seps=[53, 106])}</div></div>'
                   for _ in range(1))
    s14 = sec(14, "틀린 단어 다시 쓰기", "B쪽에서 틀린 단어를 왼쪽 칸에 쓰고 세 번 더 써요. 다 맞았으면 좋아하는 단어로!", redo)
    return f"""
<section class="page pc-page" style="--c:{r['color']}">{header(r, 'C')}{s10}{s11}{s12}{s13}{s14}{lesson_footer(r, 'C')}
</section>"""


def review_pages(set_name, rv):
    col = SETS[set_name][1]
    srows = rv["rows"]
    hd = f"""
  <div class="hd">
    <div class="tag"><div class="set">{set_name}</div><div class="no" style="font-size:16pt">복습</div></div>
    <div class="title"><div class="sound" style="font-size:19pt">{set_name} 총복습 — {html.escape(SETS[set_name][0])}</div>
      <div class="rule">{srows[0]['page']}–{srows[-1]['page']}쪽에서 배운 소리를 섞어서 확인해요.</div></div>
    <div class="meta"><div class="ln">날짜&nbsp;&nbsp;&nbsp;&nbsp;/</div><div>오늘의 별</div><div class="stars">★★★</div></div>
  </div>"""
    mixed = "".join(
        f'<div class="w" style="--c:{r["color"]}"><div class="t en">{render(w, spans(r["page"], w))}</div>'
        f'<div class="lab">{r["page"]}쪽</div></div>' for r, w in rv["mixed"])
    s1 = sec(1, "섞어 읽기", "여러 소리가 섞여 있어요. 막히는 단어는 그 쪽으로 돌아가 팁을 다시 봐요.",
             f'<div class="read r5">{mixed}</div>')
    ch = []
    for i, (r, w, (s0, e0), opts) in enumerate(rv["choice"], 1):
        shown = html.escape(w[:s0]) + '<span class="blank sm"></span>' + html.escape(w[e0:])
        gl = f'<span class="gl">({html.escape(GLOSS[w.lower()])})</span>' if w.lower() in GLOSS else ""
        ch.append(f'<div class="cho"><span class="num">{i}.</span><span class="en">{shown}</span>{gl}'
                  f'<span class="opts en">{html.escape(opts[0])}&nbsp;/&nbsp;{html.escape(opts[1])}</span></div>')
    s2 = sec(2, "알맞은 철자 고르기", "뜻을 힌트로, 두 철자 중 맞는 것에 ○ 하고 빈칸에 써요.", f'<div class="choice">{"".join(ch)}</div>')
    lines = "".join(f'<div class="d">{i}.</div>' for i in range(1, 9))
    s3 = sec(3, "받아쓰기", "단어 8개를 듣고 써요. (불러 줄 단어는 정답지에)", f'<div class="dict">{lines}</div>')
    p1 = f'<section class="page" style="--c:{col}">{hd}{s1}{s2}{s3}{footer(f"Phonics Workbook · {set_name} 총복습", "복습 1/2")}</section>'

    grid = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>" for row in rv["grid"])
    wl = "".join(f'<div class="wsw en">☐ {html.escape(w)}</div>' for w, _ in rv["placed"])
    s4 = sec(4, "낱말 찾기", "숨은 단어를 찾아 동그라미 해요. (→ 오른쪽, ↓ 아래쪽)",
             f'<div class="ws"><table class="wsg en">{grid}</table><div class="wsl">{wl}</div></div>')
    trs = "".join(
        f'<tr><td>{r["page"]}</td><td class="en" style="color:{r["color"]};font-weight:700">{html.escape(r["sound"])}</td>'
        f'<td>😀 🙂 😐</td><td>😀 🙂 😐</td><td></td></tr>' for r in srows)
    s5 = sec(5, "나의 실력 체크", "어려운 쪽은 '다시 볼 날'을 적고 한 번 더 공부해요.",
             f'<table class="self-t"><tr><th>쪽</th><th>소리</th><th>읽기</th><th>쓰기</th><th>다시 볼 날</th></tr>{trs}</table>')
    p2 = f'<section class="page" style="--c:{col}">{hd}{s4}{s5}{footer(f"Phonics Workbook · {set_name} 총복습", "복습 2/2")}</section>'
    return p1 + p2


def cover(rows):
    tracker = []
    for s, (_, col) in SETS.items():
        dots = "".join(f'<span class="dot" style="border-color:{col}"><b>{r["page"]}</b><i>A B C</i></span>'
                       for r in rows if r["set"] == s)
        tracker.append(f'<div class="rowset"><span class="lbl" style="color:{col}">{s}</span>{dots}'
                       f'<span class="dot rvd" style="border-color:{col};color:{col}">복습</span></div>')
    n_words = sum(len(r["read"]) for r in rows)
    return f"""
<section class="page cover">
  <h1 class="en">Phonics Workbook</h1>
  <div class="sub">파닉스 워크북 · 5 Sets · 42 Lessons · {n_words} Words</div>
  <div class="letters en"><span style="color:#E4572E">a</span> <span style="color:#2E86AB">sh</span> <span style="color:#3BA55C">ee</span> <span style="color:#8E5BD0">ar</span> <span style="color:#E09F1F">-tion</span></div>
  <div class="name">이름 :</div>
  <div class="name">시작한 날 :</div>
  <div class="tracker"><h3>진도표 — 끝낸 A·B·C에 색칠하거나 스티커를 붙여요</h3>{''.join(tracker)}</div>
</section>"""


def guide():
    return """
<section class="page guide" style="--c:#1d2433">
  <h1 class="pt">이렇게 공부해요</h1>
  <div class="gbox"><h3>하루 10~15분, 한 소리를 사흘에 걸쳐</h3>
  <table class="gt">
    <tr><th>날</th><th>쪽</th><th>활동</th><th>왜 하나요?</th></tr>
    <tr><td>1일차</td><td><b>A 배우기</b></td><td>💡 소리 팁 → ① 읽기 → ② 블렌딩 → ③ 찾기·분류 → ④ 복습</td><td>소리를 귀와 입으로 익히고 글자와 연결해요.</td></tr>
    <tr><td>2일차</td><td><b>B 연습</b></td><td>⑤ 따라 쓰기 → ⑥ 빈칸 → ⑦ 뜻 → ⑧ 문장 → ⑨ 받아쓰기</td><td>소리를 철자로 직접 써 보면 더 오래 기억해요.</td></tr>
    <tr><td>3일차</td><td><b>C 쓰기</b></td><td>⑩ 그림 보고 쓰기 → ⑪ 보고 쓰기 → ⑫ 문장 따라·옮겨 쓰기 → ⑬ 나만의 문장 → ⑭ 틀린 단어 다시 쓰기</td><td>단어에서 문장까지 손으로 써서 철자를 굳혀요.</td></tr>
    <tr><td>Set 끝</td><td><b>총복습</b></td><td>섞어 읽기 · 철자 고르기 · 받아쓰기 · 낱말 찾기 · 실력 체크</td><td>섞어서 꺼내 보는 연습이 진짜 실력이 돼요.</td></tr>
  </table></div>

  <div class="gbox"><h3>활동별 진행 팁 (어른용)</h3>
  <ul>
    <li><b>💡 소리 팁</b> — 입 모양을 거울로 같이 보며 소리를 3번 따라 해요.</li>
    <li><b>① 읽기</b> — 색깔 글자를 먼저 소리 내고 단어 전체를 읽어요. 한 번 읽을 때마다 ○ 하나. 속도보다 정확도가 먼저예요.</li>
    <li><b>② 블렌딩</b> — 칸 하나에 소리 하나. 손가락으로 칸을 짚으며 /c/ /a/ /t/ → cat. 회색 칸은 소리 없는 글자예요.</li>
    <li><b>③ 찾기·분류</b> — 같은 소리라도 철자가 다를 수 있어요. 눈으로 비교하며 분류하고 정답지로 확인해요.</li>
    <li><b>④ 복습</b> — 1·3·7쪽 전 단어가 다시 나와요(간격 반복). 막히면 그 쪽으로 돌아가요.</li>
    <li><b>⑤ 따라 쓰기</b> — 빨간 선은 글자가 앉는 줄, 점선은 소문자 높이예요. 쓰면서 소리를 작게 말해요.</li>
    <li><b>⑦ 뜻 · ⑧ 문장</b> — 읽은 단어의 뜻을 알면 기억이 단단해져요. 문장은 손가락으로 짚으며 두 번 읽어요.</li>
    <li><b>⑨ 받아쓰기</b> — 정답지의 단어와 문장을 자연스러운 속도로 두 번 불러 줘요. 혼자 할 때는 QR을 찍어 들어요. 틀린 단어는 '자기 점검' 칸에 다시 써요.</li>
    <li><b>⑩~⑭ 쓰기</b> — 그림 → 단어 → 문장 순서로 점점 스스로 쓰는 양을 늘려요. ⑬은 정답이 없어요. 오늘 단어를 넣었으면 칭찬해 주세요. ⑭는 B쪽에서 틀린 단어를 반복해서 써요.</li>
    <li><b>QR</b> — 휴대폰 카메라로 찍으면 '단어 듣기'는 구글 번역 발음, '소리 영상'은 유튜브 파닉스 영상 검색이 열려요(인터넷 필요).</li>
  </ul></div>

  <div class="gbox"><h3>기호 안내</h3>
  <div class="sym"><span><b class="hi" style="color:#E4572E">색깔 글자</b> 오늘의 소리</span><span>○○○ 읽은 횟수</span>
  <span><span class="box" style="--c:#2E86AB"></span> 소리 글자를 쓰는 칸</span><span><span class="star">★</span> 보충 단어</span>
  <span><span class="chunk silent en">e</span> 소리 없는 글자</span></div></div>
</section>"""


def toc(rows):
    trs = []
    for r in rows:
        trs.append(
            f'<tr><td>{r["page"]}</td><td><span class="chip" style="background:{r["color"]}">{r["set"]}</span></td>'
            f'<td class="snd en" style="color:{r["color"]}">{html.escape(r["sound"])}</td><td>{html.escape(r["rule"])}</td></tr>')
        if r["page"] in SET_END:
            trs.append(f'<tr class="rvrow"><td></td><td colspan="3">↳ {r["set"]} 총복습 (2쪽)</td></tr>')
    return f"""
<section class="page"><h1 class="pt">목차 <small>각 소리는 A(배우기)·B(연습)·C(쓰기) 3쪽</small></h1>
<table class="toc"><tr><th>쪽</th><th>세트</th><th>소리</th><th>묶음 기준</th></tr>{''.join(trs)}</table></section>"""


def key(rows, reviews):
    blocks = []
    for r in rows:
        if "sort" in r:
            cols, _, ans = r["sort"]
            a3 = "③ 분류 — " + " · ".join(
                f"<b>{html.escape(c)}</b>: <span class='en'>{html.escape(', '.join(ans[c]) or '-')}</span>" for c in cols)
        else:
            a3 = "③ ○ — <span class='en'>" + html.escape(", ".join(r["find"][1])) + "</span> (나머지 ✕)"
        a7 = "⑦ " + html.escape(", ".join(f"{w} = {GLOSS[w.lower()]}" for w in r["match"][0]))
        a8 = "⑧ <span class='en'>" + html.escape(", ".join(a for _, a in r["sent"][0])) + "</span>"
        a8 += " &nbsp; ⑩ <span class='en'>" + html.escape(", ".join(r["pictures"])) + "</span>"
        a9 = ("⑥·⑨ 단어 — <b class='en'>" + html.escape(", ".join(r["dict"][:5])) + "</b>"
              "<br>⑨ 문장 — <b class='en'>" + html.escape(r["dict_sentence"]) + "</b>")
        blocks.append(
            f'<tr><td class="kp" style="color:{r["color"]}">{r["page"]}</td><td>{a9}<br>{a3}<br>{a7}<br>{a8}</td></tr>')
        if r["page"] in SET_END:
            rv = reviews[r["set"]]
            cho = ", ".join(f'{i}. {w}' for i, (_, w, _, _) in enumerate(rv["choice"], 1))
            blocks.append(
                f'<tr class="rvrow"><td class="kp">복습</td><td><b>{r["set"]} 총복습</b><br>'
                f'② 철자 — <span class="en">{html.escape(cho)}</span><br>'
                f'③ 받아쓰기 — <b class="en">{html.escape(", ".join(rv["dict"]))}</b><br>'
                f'④ 낱말 찾기 — <span class="en">{html.escape(", ".join(w for w, _ in rv["placed"]))}</span></td></tr>')
    pages, per = [], 8
    for i in range(0, len(blocks), per):
        pages.append(f"""
<section class="page"><h1 class="pt">정답 · 불러 주기 ({i // per + 1})</h1>
<table class="key">{''.join(blocks[i:i + per])}</table></section>""")
    grids = []
    for s, rv in reviews.items():
        g = "".join("<tr>" + "".join(
            f'<td class="{"on" if (a, b) in rv["sol"] else ""}">{c}</td>' for b, c in enumerate(row)) + "</tr>"
            for a, row in enumerate(rv["grid"]))
        grids.append(f'<div class="kg" style="--c:{SETS[s][1]}"><div class="kgt">{s}</div><table class="wsg sm en">{g}</table></div>')
    pages.append(f"""
<section class="page"><h1 class="pt">정답 · 낱말 찾기</h1><div class="kgs">{''.join(grids)}</div></section>""")
    return "".join(pages)


def build():
    rows = build_lesson_data(load_rows())
    reviews = build_reviews(rows)
    body = [cover(rows), guide(), toc(rows)]
    for r in rows:
        body += [page_a(r), page_b(r), page_c(r)]
        if r["page"] in SET_END:
            body.append(review_pages(r["set"], reviews[r["set"]]))
    body.append(key(rows, reviews))
    with open(CSS_FILE, encoding="utf-8") as f:
        css = f.read()
    doc = f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>Phonics Workbook</title>
<link href="https://fonts.googleapis.com/css2?family=Andika:wght@400;700&family=Noto+Sans+KR:wght@400;700;900&display=swap" rel="stylesheet">
<style>{css}</style></head><body>{''.join(body)}</body></html>"""
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(doc)
    subprocess.run([
        CHROME, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
        "--virtual-time-budget=20000", f"--print-to-pdf={OUT_PDF}", "file://" + OUT_HTML,
    ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("wrote", OUT_HTML, OUT_PDF, file=sys.stderr)


if __name__ == "__main__":
    build()
